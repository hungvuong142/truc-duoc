"""Builds the monthly "Bảng phân công trực" Word document (see
templates/word_template.docx for the target layout) from RosterRow data
(app/logic/roster_export.py). Pure python-docx construction -- no DB access,
no Streamlit -- so app/ui/outputs_page.py just fetches data, builds
RosterRows, and hands them to build_roster_document().

Layout constants below are the single place to tweak fonts/margins/text if
the hospital's letterhead or signature block ever changes.
"""

from __future__ import annotations

from datetime import datetime

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.config import BASE_HANOI, BASE_NINH_BINH
from app.logic.roster_export import RosterRow

# --------------------------------------------------------------------------
# Page / font
# --------------------------------------------------------------------------
# Config here
FONT_NAME = "Times New Roman"
# Config here
FONT_SIZE_PT = 12
# Config here
MARGINS_CM = {"top": 2, "bottom": 2, "left": 1.5, "right": 1}

# --------------------------------------------------------------------------
# Header (2x1, borders=none)
# --------------------------------------------------------------------------
# Config here
HEADER_COL_WIDTHS_CM = [6.94, 11.56]
# Config here
HEADER_LINE_SPACING = 1.05
# Config here
HEADER_LEFT_LINES = ["BỆNH VIỆN BẠCH MAI", "KHOA DƯỢC"]
# Config here
HEADER_RIGHT_LINES = [
    "Hotline Khoa Dược CSHN: 086 9587698",
    "Hotline Khoa Dược CSNB: 086 9572788",
]

# Config here
DOC_TITLE_TEMPLATE = "BẢNG PHÂN CÔNG TRỰC KHOA DƯỢC THÁNG {month:02d} NĂM {year:04d}"

# --------------------------------------------------------------------------
# Roster tables
# --------------------------------------------------------------------------
# Config here
SECTION_HEADINGS = {
    BASE_HANOI: "CƠ SỞ HÀ NỘI",
    BASE_NINH_BINH: "CƠ SỞ NINH BÌNH",
}
# Config here -- smaller than the document's base font: long full names
# (Đại học/Cao đẳng columns) were wrapping onto multiple lines at 12pt; a
# smaller table font plus the wider name columns below give names more
# room to fit on one line.
TABLE_FONT_SIZE_PT = 10
# Config here
TABLE_COL_WIDTHS_CM = [1.5, 1.2, 3.7, 3.7, 3.7, 3.7, 1.0]
# Config here -- Word's "Allow row to break across pages", off, so a tall
# multi-name row is pushed whole onto the next page instead of being cut
# mid-row.
PREVENT_TABLE_ROW_SPLIT = True
# Config here -- the 2-row header (Ngày/Thứ/DS Đại học/DS Cao đẳng/Ghi chú +
# Trực 24/24 vs 12/24) repeats on every page the table spans, so a page
# break mid-month still shows which column is which.
ROSTER_HEADER_ROW_COUNT = 2

# --------------------------------------------------------------------------
# Signature block (2x1, borders=none)
# --------------------------------------------------------------------------
# Config here
SIGNATURE_COL_WIDTHS_CM = [9.25, 9.25]
# Config here
SIGNATURE_PLACE = "Hà Nội"
# Config here
SIGNATURE_TITLE_LINE = "TRƯỞNG KHOA DƯỢC"
# Config here
SIGNATURE_NAME = "Nguyễn Thu Minh"
# Config here
SIGNATURE_BLANK_LINES_BEFORE_NAME = 4


def _set_run_font(run, bold: bool = False, italic: bool = False, size: int | None = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size or FONT_SIZE_PT)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT_NAME)


def _force_rfonts(rPr, font_name: str) -> None:
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font_name)


def _set_all_styles_font(doc: Document) -> None:
    """Force Times New Roman on every style, not just Normal -- Word/
    LibreOffice otherwise fall back to a theme font (usually Calibri) for
    styles (e.g. Table Grid) that don't inherit font from Normal."""
    for style in doc.styles:
        try:
            font = style.font
        except (AttributeError, NotImplementedError):
            continue
        if font is None:
            continue
        font.name = FONT_NAME
        try:
            rPr = style.element.get_or_add_rPr()
            _force_rfonts(rPr, FONT_NAME)
        except Exception:
            pass

    styles_element = doc.styles.element
    doc_defaults = styles_element.find(qn("w:docDefaults"))
    if doc_defaults is not None:
        rpr_default = doc_defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                rpr_default.append(rpr)
            _force_rfonts(rpr, FONT_NAME)


def _remove_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _set_para_spacing(paragraph, line_spacing: float, before_pt: float = 0, after_pt: float = 0) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)


def _prevent_row_split(table) -> None:
    """Sets cantSplit on every row of `table` -- python-docx has no
    high-level API for this, so it has to be done directly in the row's
    XML (<w:trPr><w:cantSplit/></w:trPr>)."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))


def _set_repeat_header_rows(table, num_rows: int) -> None:
    """Marks the table's first `num_rows` rows as a header that Word repeats
    on every page the table spans -- python-docx has no high-level API for
    this, so <w:tblHeader/> is set directly in each row's trPr."""
    for row in table.rows[:num_rows]:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))


def _set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for col, width_cm in zip(table.columns, widths_cm):
        col.width = Cm(width_cm)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])


def set_document_defaults(doc: Document) -> None:
    _set_all_styles_font(doc)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)
    _force_rfonts(style.element.get_or_add_rPr(), FONT_NAME)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(MARGINS_CM["top"])
    section.bottom_margin = Cm(MARGINS_CM["bottom"])
    section.left_margin = Cm(MARGINS_CM["left"])
    section.right_margin = Cm(MARGINS_CM["right"])


def add_header(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)
    _set_col_widths(table, HEADER_COL_WIDTHS_CM)

    left_cell, right_cell = table.rows[0].cells

    for i, text in enumerate(HEADER_LEFT_LINES):
        p = left_cell.paragraphs[0] if i == 0 else left_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_para_spacing(p, HEADER_LINE_SPACING)
        run = p.add_run(text)
        _set_run_font(run, bold=(i == 1))

    for i, text in enumerate(HEADER_RIGHT_LINES):
        p = right_cell.paragraphs[0] if i == 0 else right_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_para_spacing(p, HEADER_LINE_SPACING)
        run = p.add_run(text)
        _set_run_font(run, italic=True)


def add_title(doc: Document, month: int, year: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DOC_TITLE_TEMPLATE.format(month=month, year=year))
    _set_run_font(run, bold=True)


def _add_multiline_run(paragraph, text: str, bold: bool = False, italic: bool = False, size: int | None = None) -> None:
    """Adds `text` to `paragraph`, turning any "\\n" into a real line break
    (a bare "\\n" inside a run's <w:t> does not render as a line break in
    Word) so multi-name cells (see roster_export.NAME_SEPARATOR) show one
    name per line instead of one unbroken run of text."""
    last_run = None
    for i, line in enumerate(text.split("\n")):
        if last_run is not None:
            last_run.add_break()
        last_run = paragraph.add_run(line)
        _set_run_font(last_run, bold=bold, italic=italic, size=size)


def _add_centered_run(cell, text: str, bold: bool = False, size: int | None = None) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, bold=bold, size=size)


def _add_roster_header_rows(table) -> None:
    header1 = table.rows[0].cells
    header2 = table.rows[1].cells

    _add_centered_run(header1[0], "Ngày", bold=True, size=TABLE_FONT_SIZE_PT)
    header1[0].merge(header2[0])

    _add_centered_run(header1[1], "Thứ", bold=True, size=TABLE_FONT_SIZE_PT)
    header1[1].merge(header2[1])

    _add_centered_run(header1[2], "DS Đại học", bold=True, size=TABLE_FONT_SIZE_PT)
    header1[2].merge(header1[3])

    _add_centered_run(header1[4], "DS Cao đẳng", bold=True, size=TABLE_FONT_SIZE_PT)
    header1[4].merge(header1[5])

    _add_centered_run(header1[6], "Ghi chú", bold=True, size=TABLE_FONT_SIZE_PT)
    header1[6].merge(header2[6])

    for idx in (2, 3, 4, 5):
        label = "Trực 24/24" if idx in (2, 4) else "Trực 12/24"
        _add_centered_run(header2[idx], label, size=TABLE_FONT_SIZE_PT)


def add_base_section(doc: Document, heading: str, rows: list[RosterRow]) -> None:
    heading_p = doc.add_paragraph()
    heading_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(heading_p.add_run(heading), bold=True)

    table = doc.add_table(rows=2, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _add_roster_header_rows(table)

    for row in rows:
        cells = table.add_row().cells
        values = [
            f"{row.day.day}/{row.day.month}",
            row.weekday_label,
            row.dai_hoc_24,
            row.dai_hoc_12,
            row.cao_dang_24,
            row.cao_dang_12,
            row.ghi_chu,
        ]
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            _add_multiline_run(p, value, size=TABLE_FONT_SIZE_PT)

    _set_col_widths(table, TABLE_COL_WIDTHS_CM)
    if PREVENT_TABLE_ROW_SPLIT:
        _prevent_row_split(table)
    _set_repeat_header_rows(table, ROSTER_HEADER_ROW_COUNT)


def add_signature(doc: Document, generated_at: datetime) -> None:
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)
    _set_col_widths(table, SIGNATURE_COL_WIDTHS_CM)
    _, right_cell = table.rows[0].cells

    date_p = right_cell.paragraphs[0]
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(date_p, HEADER_LINE_SPACING)
    date_line = f"{SIGNATURE_PLACE}, ngày {generated_at.day:02d} tháng {generated_at.month:02d} năm {generated_at.year:04d}"
    _set_run_font(date_p.add_run(date_line))

    title_p = right_cell.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(title_p, HEADER_LINE_SPACING)
    _set_run_font(title_p.add_run(SIGNATURE_TITLE_LINE), bold=True)

    for _ in range(SIGNATURE_BLANK_LINES_BEFORE_NAME):
        _set_para_spacing(right_cell.add_paragraph(), HEADER_LINE_SPACING)

    name_p = right_cell.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(name_p, HEADER_LINE_SPACING)
    _set_run_font(name_p.add_run(SIGNATURE_NAME), bold=True)


def build_roster_document(
    month: int,
    year: int,
    tables: dict[str, list[RosterRow]],
    generated_at: datetime,
) -> Document:
    doc = Document()
    set_document_defaults(doc)

    add_header(doc)
    doc.add_paragraph()
    add_title(doc, month, year)

    for i, base in enumerate((BASE_HANOI, BASE_NINH_BINH)):
        if i > 0:
            doc.add_paragraph()
        add_base_section(doc, SECTION_HEADINGS[base], tables.get(base, []))

    add_signature(doc, generated_at)
    return doc
